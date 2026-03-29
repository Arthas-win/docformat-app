import os
import re
from django.conf import settings
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.parser import OxmlElement
from docx.oxml.ns import qn


def format_document(input_path, output_path, formatting_options):
    """
    Apply formatting options to a DOCX document
    
    Args:
        input_path (str): Path to the input DOCX file
        output_path (str): Path where the formatted DOCX file will be saved
        formatting_options (dict): Dictionary containing formatting options
    """
    # Load the input document
    doc = Document(input_path)
    
    # Extract formatting options with defaults
    font_family = formatting_options.get('fontFamily', 'Times New Roman')
    font_size = int(formatting_options.get('fontSize', 14)) # Зазвичай 14 для академічних робіт
    line_spacing = float(formatting_options.get('lineSpacing', 1.5))
    add_page_numbers = formatting_options.get('pageNumbers', True)
    
    # Process all paragraphs in the document
    for paragraph in doc.paragraphs:
        # Apply font and size to runs in the paragraph
        for run in paragraph.runs:
            run.font.name = font_family
            run.font.size = Pt(font_size)
        
        # Apply line spacing and paragraph formatting
        paragraph_format = paragraph.paragraph_format
        paragraph_format.line_spacing = line_spacing
        # Set spacing before and after paragraphs to standard academic format
        paragraph_format.space_before = Pt(0)
        paragraph_format.space_after = Pt(0)
        
        # Standard paragraph indent (абзацний відступ) 1.25 cm
        paragraph_format.first_line_indent = Inches(0.49) # 1.25 cm
        
        # Set alignment (justified for body text according to academic standards)
        paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    # Process tables if any exist
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = font_family
                        run.font.size = Pt(max(10, font_size - 2)) # В таблицях шрифт зазвичай на 2пт менше (12 якщо основний 14)
                    
                    # Таблиці мають мати одинарний інтервал і без абзацного відступу
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.paragraph_format.first_line_indent = Pt(0)
                    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    # Process headers and footers in sections
    for section in doc.sections:
        # Apply section formatting (margins per Ukrainian academic standards)
        from docx.shared import Mm
        section.top_margin = Mm(20)      # Top margin: 20 mm
        section.bottom_margin = Mm(20)   # Bottom margin: 20 mm
        section.left_margin = Mm(30)     # Left margin: 30 mm
        section.right_margin = Mm(15)    # Right margin: 15 mm
        
        # Add page numbers if requested
        if add_page_numbers:
            add_page_numbers_to_section(section, doc)
    
    # Update default paragraph style
    update_default_paragraph_style(doc.styles, font_family, font_size, line_spacing)
    
    # Save the formatted document
    doc.save(output_path)


def update_default_paragraph_style(styles, font_family, font_size, line_spacing):
    """
    Update the default paragraph style in the document
    """
    try:
        # Get the default paragraph style
        paragraph_style = styles['Normal']
        if paragraph_style:
            paragraph_style.font.name = font_family
            paragraph_style.font.size = Pt(font_size)
            paragraph_style.paragraph_format.line_spacing = line_spacing
    except KeyError:
        # Style doesn't exist, continue without modification
        pass


def add_page_numbers_to_section(section, doc):
    """
    Add page numbers to the section's footer
    """
    footer = section.footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Add page number field that will be calculated by Word
    page_num_run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    
    page_num_run._r.append(fldChar1)
    page_num_run._r.append(instrText)
    page_num_run._r.append(fldChar2)


def format_document_service(job):
    """
    Main service function to process a document formatting job
    
    Args:
        job (DocumentJob): The document job model instance
        
    Returns:
        str: Path to the output file
    """
    # Get the input file path
    input_file_path = job.input_file.path
    
    # Define output path
    output_dir = os.path.join(settings.MEDIA_ROOT, "jobs", "output")
    os.makedirs(output_dir, exist_ok=True)
    
    # Generate output filename based on job ID
    output_filename = f"formatted_{job.id}.docx"
    output_path = os.path.join(output_dir, output_filename)
    
    # Get formatting options from job metadata
    formatting_options = job.meta_json or {}
    
    # Apply formatting
    format_document(input_file_path, output_path, formatting_options)
    
    # Update job progress
    job.progress = 100
    job.current_stage = "completed"
    job.status = job.Status.DONE
    
    # Save the output file to the job model
    with open(output_path, 'rb') as output_file:
        from django.core.files import File
        job.output_docx.save(
            os.path.basename(output_path),
            File(output_file),
            save=True
        )
    
    return output_path