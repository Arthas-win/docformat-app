from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


def _paragraph_has_page_break(paragraph):
    for br in paragraph._p.xpath(".//w:br"):
        if br.get(qn("w:type")) == "page":
            return True
    return False


def _append_page_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "

    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_separate)
    run._r.append(fld_end)


def _set_page_numbers(doc, enabled):
    for section in doc.sections:
        section.different_first_page_header_footer = True
        footer = section.footer
        paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        paragraph.clear()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if enabled:
            _append_page_field(paragraph)


def format_docx_except_first_page(
    file_obj,
    font_family="Times New Roman",
    font_size=14,
    line_spacing=1.5,
    page_numbers=True,
):
    document = Document(file_obj)

    first_page = True
    for paragraph in document.paragraphs:
        if not first_page:
            paragraph.paragraph_format.line_spacing = line_spacing
            for run in paragraph.runs:
                if run.text:
                    run.font.name = font_family
                    run.font.size = Pt(font_size)
        if _paragraph_has_page_break(paragraph):
            first_page = False

    _set_page_numbers(document, page_numbers)

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output
